#!/usr/bin/env python3
"""Build the full redesigned MRM deck (slides 5-38) from the untouched master."""
import shutil, os, sys, copy
from deck import *
import deck as D
from slides import *
import mom_data
import rag as R
from rag import FILL, DARK

SRC="master.pptx"
OUT="MRM_Aug-26_Redesigned_V8_RAG.pptx"
XLS="MRM_Master_Data.xlsx"

# ── auto row heights for the MOM tables ─────────────────────────────────────
def mom_heights(rowsd, wtotal=RIGHT-M, avail=None, size=T_TBLC):
    """Row heights proportional to the tallest cell, at a given cell font size."""
    colw=[c/sum(MOM_COLW)*wtotal for c in MOM_COLW]
    cpi = 1.0/(size*0.0092)          # characters per inch at `size` on a 1.5x canvas
    hs=[]
    for row in rowsd:
        lines=1
        for i,cell in enumerate(row):
            usable=colw[i]-0.24
            n=0
            for seg in str(cell).split("\n"):
                n+=max(1,-(-len(seg)//max(4,int(usable*cpi))))
            lines=max(lines,n)
        hs.append(lines*(size*1.30/72.0)+0.26)
    if avail and sum(hs)<avail:
        extra=(avail-sum(hs))/len(hs)
        hs=[h+extra for h in hs]
    return hs

def mom_fit(rowsd, avail):
    """Largest cell font at which every row fits without clipping."""
    for size in range(T_TBLC, 11, -1):
        hs=mom_heights(rowsd, size=size)
        if sum(hs)<=avail:
            return size, mom_heights(rowsd, avail=avail, size=size)
    return 12, mom_heights(rowsd, avail=avail, size=12)

def mom(sp, page, rowsd):
    title_block(sp,"Update on Last MOM",
                f"Minutes of meeting · action tracker   ·   page {page} of 3")
    avail=BODY_BOT-BODY_TOP-0.50
    body=[]; chips={}; cellc={}
    LBL={'g':"Closed",'a':"WIP",'r':"Open"}
    for r,row in enumerate(rowsd):
        st=row[6]
        R.note(f"MOM p{page} action {r+1} · {row[1].splitlines()[0]}",
               "action status from the TPL update", LBL[st], st)
        body.append(list(row[:6])+[LBL[st]])
        chips[(r,6)]=st; cellc[(r,0)]=DARK[st]
    size,hs=mom_fit(body, avail-0.34)
    table(sp,M,BODY_TOP,RIGHT-M,MOM_HDR,body,MOM_COLW,
          rh=lambda r: hs[r], hh=0.50, cell=size, hdr=T_TBLH,
          aligns=[PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.LEFT,PP_ALIGN.LEFT,
                  PP_ALIGN.LEFT,PP_ALIGN.LEFT,PP_ALIGN.CENTER],
          chips=chips, cell_colors=cellc)
    rag_legend(sp)

# ── photo slide specs: slide -> (title, subtitle, captions, cols) ───────────
PHOTO={
 21:("Good Practices","Environment, health and safety   ·   August 2026",
     ["Illumination inspection at the work front","Noise level monitoring — 54.5 dB",
      "Noise level monitoring — 54.0 dB","Pedestrian access board displayed near Gate No. 02",
      "Illumination inspection — night shift"],3),
 22:("HSE Good Practices","Vehicle safety and recognition   ·   August 2026",
     ["Seven Mirror Concept implemented on all heavy vehicles",
      "Safety Conscious Person of the Month displayed"],2),
 23:("Environmental Initiative","Signages produced from scrap plywood   ·   August 2026",
     ["Deep excavation caution board","Deep excavation caution board — work front",
      "Speed limit and PPE zone signage","Site safety notice board",
      "Lifting safety instruction board","Safety and welfare signage board"],3),
 24:("Tool Box Talks","Daily briefings delivered on site   ·   August 2026",
     ["Emergency preparedness and response — 21 Aug 26",
      "Right way of using a full body harness — 11 Jul 26",
      "Hazards in bar bending and cutting activity — 07 Aug 26",
      "Tool box talk at the work front — 11 Aug 26",
      "Site safety rules — 27 Aug 26",
      "Full workforce briefing before shift start"],3),
 25:("Safety Posters and Signages","Site communication   ·   August 2026",
     ["Safety briefings across the workforce","Child labour prohibited — statutory notice",
      "Workmen accommodation guidelines board",
      "Behavioural safety mural at the accommodation block",
      "Speed limit and work-in-progress signage with blinking lights near the concrete plant"],3),
 26:("Mock Drill","Worker hit by a moving vehicle   ·   emergency response rehearsal",
     ["Emergency response, casualty handling and evacuation sequence rehearsed on site"],1),
 29:("Quality Good Practices","Workmanship standards adopted on site   ·   August 2026",
     ["Steel coupler threading in progress","Go / No-Go gauge test after coupler threading",
      "Curing compound applied uniformly after pouring","Curing compound coverage across the pour area",
      "Sample installation of Formstop to avoid column starters",
      "Formstop fixing at the column starter location"],3),
 30:("Training Imparted","Quality and skill development   ·   August 2026",
     ["Steel cutting and bending training","Steel cutting and banding training",
      "Steel threading training","Threading machine demonstration",
      "Coupler threading quality check","Concrete pouring and finishing training",
      "Site quality induction","Site quality induction — classroom session",
      "Tool box training at the work front","Workmen briefing at the work front"],5),
 31:("4 Star Hotel Progress","July 2026 versus August 2026",
     ["July 2026","August 2026"],2),
 32:("3 Star Hotel Progress","July 2026 versus August 2026",
     ["July 2026","August 2026"],2),
 33:("Convention Centre Progress","July 2026 versus August 2026",
     ["July 2026","August 2026"],2),
 35:("Progress Photos","Foundation completion by area   ·   August 2026",
     ["4 Star area — 39 of 46 footings completed","4 Star area — column and pedestal casting",
      "3 Star area — 15 of 51 footings completed","3 Star area — raft and column reinforcement",
      "Convention Centre — 08 of 50 footings completed"],3),
 36:("Site Infrastructure","Enabling works and welfare   ·   August 2026",
     ["Safety park and information kiosk","Workmen rest area and pedestrian walkway",
      "Tower crane — operational","Site toilet facility for workmen"],2),
 37:("Workmen Camp & Steel Yard","Camp delivery and material handling   ·   August 2026",
     ["Camp development and toilet shed","Workmen shed C1 — G+2",
      "Medical and EHS room","Steel yard"],2),
}

def main():
    if os.path.exists(OUT): os.remove(OUT)
    shutil.copyfile(SRC,OUT)
    prs=Presentation(OUT)
    S=prs.slides
    D.scan_chrome(prs)

    def slide(n): return S[n-1]

    # ---- 5 .. 20 -----------------------------------------------------------
    simple={5:s05,6:s06,7:s07,8:s08,9:s09,10:s10,11:s11,12:s12,13:s13,
            17:s17,19:s19,20:s20,27:s27}
    for n,fn in simple.items():
        sl=slide(n); clear_slide(sl); fn(sl.shapes)

    for n,page,data in ((14,1,mom_data.S14),(15,2,mom_data.S15),(16,3,mom_data.S16)):
        sl=slide(n); clear_slide(sl); mom(sl.shapes,page,data)

    # ---- 18: keep the programme artwork / OLE, restyle the chrome ----------
    sl=slide(18)
    kept=clear_slide(sl, keep_pictures=True, keep_ole=True)
    title_block(sl.shapes,"September 2026 Pour Plan",
                "Block-wise pour plan and RCC quantities   ·   September 2026")
    E=914400
    art=[sh for sh in kept if sh.width and sh.width/E>2.0 and sh.height/E>2.0]
    small=[sh for sh in kept if sh not in art and not is_logo(sh)]
    for sh in small:
        sh._element.getparent().remove(sh._element)
    art.sort(key=lambda sh:-(sh.width*sh.height))
    ch=BODY_BOT-BODY_TOP
    cells=[]
    if len(art)==1:
        cells=[(M,BODY_TOP,RIGHT-M,ch)]
    else:
        cells=[(M,BODY_TOP,11.90,ch)]
        rest=len(art)-1
        rh=(ch-(rest-1)*GUT)/rest
        for i in range(rest):
            cells.append((M+11.90+GUT,BODY_TOP+i*(rh+GUT),RIGHT-M-11.90-GUT,rh))
    for sh,(cx,cy,cwid,chh) in zip(art,cells):
        ar=sh.width/sh.height
        w=min(cwid,chh*ar); h=w/ar
        sh.left=Inches(cx+(cwid-w)/2); sh.top=Inches(cy+(chh-h)/2)
        sh.width=Inches(w); sh.height=Inches(h)

    # ---- 28: quality performance (charts rebuilt, originals dropped) -------
    sl=slide(28); clear_slide(sl); s28(sl.shapes,None)

    # ---- photo slides ------------------------------------------------------
    for n,(t,sub,caps,cols) in PHOTO.items():
        sl=slide(n)
        imgs=harvest_images(sl)
        clear_slide(sl)
        title_block(sl.shapes,t,sub)
        if imgs:
            photo_slide(sl.shapes,imgs,caps,cols=cols)
        else:
            text(tbox(sl.shapes,M,5.0,RIGHT-M,0.5),"Photographs to be inserted",
                 T_BODY,MUTED,align=PP_ALIGN.CENTER)
        print(f"  slide {n}: {len(imgs)} images")

    # ---- 34 divider, 38 closing -------------------------------------------
    sl=slide(34); clear_slide(sl)
    divider(sl.shapes,"Updates","Progress · Quality · Photographs")
    sl=slide(38); clear_slide(sl)
    divider(sl.shapes,"Thank You","Adani Lucknow Airport City   ·   Monthly Review   ·   August 2026")

    # original slide 9 was flagged hidden; the rebuilt clash report must be visible
    for n in range(5,39):
        S[n-1]._element.set('show','1')

    # ---- one worksheet per chart, written from the charts themselves -------
    D.write_workbook(XLS, R.ROWS)
    print("  workbook sheets:", len(D.CHARTS)+2, " RAG rules:", len(R.ROWS))

    # ---- drop stale relationships left behind by the cleared originals -----
    import re as _re
    for sl in S:
        part=sl.part
        xml=part._element.xml
        used=set(_re.findall(r'r:(?:id|embed|link|pict|dm|lo|qs|cs)="([^"]+)"', xml))
        for rid in list(part.rels):
            rel=part.rels[rid]
            if rel.reltype.endswith('/slideLayout'): continue
            if rel.reltype.endswith('/notesSlide'): continue
            if rid not in used:
                part.rels.pop(rid)

    # ---- deck-wide hygiene -------------------------------------------------
    uid=5000
    for s in S:
        for el in s._element.iter(qn('p:cNvPr')):
            el.set('id',str(uid)); uid+=1

    prs.save(OUT)

    # attach the workbook to every chart part (embedded, refreshable)
    from pptx import Presentation as P2
    print("charts:",len(D.CHART_PARTS))
    p2=P2(OUT)
    blob=open(XLS,'rb').read()
    for sl in p2.slides:
        for sh in sl.shapes:
            if sh.has_chart:
                try:
                    part=sh.chart.part.chart_workbook.xlsx_part
                    if part is not None: part._blob=blob
                except Exception as e:
                    print("   wb:",e)
    p2.save(OUT)
    rename_embeddings(OUT)
    print("saved",OUT,os.path.getsize(OUT))

def rename_embeddings(path):
    """Give every embedded chart workbook its real file name, for auditability."""
    import zipfile, re as _re
    zin=zipfile.ZipFile(path); names=zin.namelist()
    ren={}
    for n in names:
        m=_re.match(r'ppt/embeddings/Microsoft_Excel_Sheet(\d*)\.xlsx$', n)
        if m:
            ren[n]="ppt/embeddings/MRM_Master_Data%s.xlsx" % (m.group(1) or "1")
    short={k.split('/')[-1]: v.split('/')[-1] for k,v in ren.items()}
    items=[(n, zin.read(n)) for n in names]
    zin.close()
    tmp=path+".tmp"
    zout=zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED)
    for n,data in items:
        if n=='[Content_Types].xml' or n.endswith('.rels'):
            t=data.decode('utf8')
            for a_,b_ in sorted(short.items(), key=lambda kv:-len(kv[0])):
                t=t.replace(a_,b_)
            data=t.encode('utf8')
        zout.writestr(ren.get(n,n), data)
    zout.close()
    os.replace(tmp, path)


if __name__=="__main__":
    main()
