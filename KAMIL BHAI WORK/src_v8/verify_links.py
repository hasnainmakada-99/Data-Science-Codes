#!/usr/bin/env python3
"""End-to-end Excel-link audit for every chart in a pptx.

For each chart part it:
  1. finds the embedded workbook the chart actually points at (via its own rels),
  2. reads every c:f formula (series name, categories, values),
  3. resolves that range inside the workbook with openpyxl,
  4. compares the resolved cell contents against the cache the chart carries,
  5. flags external links, missing workbooks, bad sheet names and stale ranges.
"""
import sys, re, zipfile, io, posixpath
from lxml import etree
import openpyxl

NS={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart',
    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'rel':'http://schemas.openxmlformats.org/package/2006/relationships'}
Q=lambda t:'{%s}%s'%(NS[t.split(':')[0]],t.split(':')[1])

CELL=re.compile(r"^(?:'([^']+)'|([^!]+))!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$")

def colnum(s):
    n=0
    for ch in s: n=n*26+ord(ch)-64
    return n

def parse_ref(f):
    m=CELL.match(f.strip())
    if not m: return None
    sheet=m.group(1) or m.group(2)
    c1,r1=colnum(m.group(3)),int(m.group(4))
    c2,r2=(colnum(m.group(5)),int(m.group(6))) if m.group(5) else (c1,r1)
    return sheet,r1,c1,r2,c2

def cache_of(node):
    """Cached points keyed by c:idx, plus the declared c:ptCount."""
    out={}
    for pt in node.iter(Q('c:pt')):
        v=pt.find(Q('c:v'))
        out[int(pt.get('idx'))]=None if v is None or v.text is None else v.text
    pc=node.find('.//'+Q('c:ptCount'))
    return out, (int(pc.get('val')) if pc is not None else len(out))

def norm(v):
    if v is None: return None
    if isinstance(v,(int,float)):
        return round(float(v),6)
    s=str(v).strip()
    try: return round(float(s),6)
    except ValueError: return s

def main(path):
    z=zipfile.ZipFile(path)
    charts=sorted(n for n in z.namelist() if re.match(r'ppt/charts/chart\d+\.xml$',n))
    # which slide owns which chart
    owner={}
    for n in z.namelist():
        m=re.match(r'ppt/slides/_rels/slide(\d+)\.xml\.rels',n)
        if not m: continue
        for t in re.findall(r'Target="([^"]+chart\d+\.xml)"', z.read(n).decode()):
            owner[posixpath.normpath(posixpath.join('ppt/slides',t))]=int(m.group(1))
    problems=[]; checked=0; refs=0
    print(f"{'chart':<22}{'slide':>6}  {'workbook':<26}{'sheet':<22}{'refs':>5}  result")
    print("-"*104)
    for c in charts:
        rels=c.replace('charts/','charts/_rels/')+'.rels'
        wbpart=None; external=[]
        if rels in z.namelist():
            root=etree.fromstring(z.read(rels))
            for rel in root:
                if rel.get('TargetMode')=='External': external.append(rel.get('Target'))
                if rel.get('Type','').endswith('/package') or 'embeddings' in (rel.get('Target') or ''):
                    wbpart=posixpath.normpath(posixpath.join('ppt/charts',rel.get('Target')))
        name=c.split('/')[-1]; sl=owner.get(c,'—')
        if external:
            problems.append(f"{name}: EXTERNAL link {external}")
            print(f"{name:<22}{sl:>6}  {'EXTERNAL':<26}{'-':<22}{0:>5}  FAIL external link")
            continue
        if wbpart is None or wbpart not in z.namelist():
            problems.append(f"{name}: no embedded workbook")
            print(f"{name:<22}{sl:>6}  {'(missing)':<26}{'-':<22}{0:>5}  FAIL no workbook")
            continue
        wb=openpyxl.load_workbook(io.BytesIO(z.read(wbpart)), data_only=False)
        cx=etree.fromstring(z.read(c))
        bad=[]; sheets=set(); nrefs=0
        for holder in cx.iter(Q('c:tx'),Q('c:cat'),Q('c:val')):
            f=holder.find('.//'+Q('c:f'))
            if f is None or not f.text: continue
            nrefs+=1
            pr=parse_ref(f.text)
            if not pr:
                bad.append(f"unparseable ref {f.text!r}"); continue
            sheet,r1,c1,r2,c2=pr; sheets.add(sheet)
            if sheet not in wb.sheetnames:
                bad.append(f"sheet '{sheet}' not in workbook"); continue
            ws=wb[sheet]
            live=[norm(ws.cell(row=r,column=col).value)
                  for r in range(r1,r2+1) for col in range(c1,c2+1)]
            cached,ptcount=cache_of(holder)
            if ptcount!=len(live):
                bad.append(f"{f.text}: range spans {len(live)} cells, chart declares {ptcount}")
                continue
            for i,a in enumerate(live):
                b=norm(cached.get(i))
                if i not in cached:
                    # a gap in the cache is only valid where the workbook cell is empty
                    if a is not None:
                        bad.append(f"{f.text}: cell {i+1} workbook={a!r} missing from chart cache")
                        break
                elif a!=b:
                    bad.append(f"{f.text}: cell {i+1} workbook={a!r} chart={b!r}")
                    break
        checked+=1; refs+=nrefs
        res="OK" if not bad else "FAIL — "+bad[0]
        print(f"{name:<22}{sl:>6}  {wbpart.split('/')[-1]:<26}{','.join(sorted(sheets))[:21]:<22}{nrefs:>5}  {res}")
        for b in bad: problems.append(f"{name}: {b}")
    print("-"*104)
    print(f"charts checked: {checked}   c:f references resolved: {refs}   problems: {len(problems)}")
    for p in problems: print("  !", p)
    return 1 if problems else 0

if __name__=="__main__":
    sys.exit(main(sys.argv[1]))
