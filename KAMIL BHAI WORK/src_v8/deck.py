#!/usr/bin/env python3
"""MRM Aug-26 — full deck rebuild to Design System v7.
Source of truth: master.pptx (never modified).  Output: a new copy.
"""
import copy, io, os, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from lxml import etree
import openpyxl

# ───────────────────────────────── palette
NAVY=RGBColor(0x17,0x37,0x5E); DEEP=RGBColor(0x0F,0x24,0x40)
TATA=RGBColor(0x14,0x79,0xB8); BLUTNT=RGBColor(0xE8,0xF1,0xF8); BLUMID=RGBColor(0xB9,0xD5,0xE9)
GREEN=RGBColor(0x2E,0x9E,0x5B); AMBER=RGBColor(0xE9,0xA2,0x1B); RED=RGBColor(0xD6,0x45,0x3D)
TEAL=RGBColor(0x0E,0x8C,0x8C)
GRNT=RGBColor(0xE6,0xF4,0xEB); AMBT=RGBColor(0xFC,0xF2,0xDD); REDT=RGBColor(0xFB,0xEA,0xE9)
BODY=RGBColor(0x4A,0x5A,0x6E); MUTED=RGBColor(0x8A,0x98,0xA8)
RULE=RGBColor(0xE2,0xE8,0xEF); WASH=RGBColor(0xF5,0xF8,0xFB); WHITE=RGBColor(0xFF,0xFF,0xFF)
AMBDK=RGBColor(0xB8,0x80,0x0F); GRNDK=RGBColor(0x1E,0x6E,0x3E); REDDK=RGBColor(0x96,0x30,0x2A)
FONT="Trebuchet MS"

# ───────────────────────────────── type scale  (v7 = v6 × ~0.9)
T_TITLE=40; T_SUB=16; T_SECT=24; T_CARD=21; T_BIG=34; T_MID=28; T_BODY=17; T_CAP=16
T_TBLH=16; T_TBLC=15

# ───────────────────────────────── grid
M=0.60; RIGHT=19.40; GUT=0.30; PAGE_W=20.0; PAGE_H=11.25
TITLE_Y=0.20; SUB_Y=0.94; RULE_Y=1.36; BODY_TOP=1.56; BODY_BOT=10.62
def I(v): return Inches(v)

CS_ORDER=['c:date1904','c:lang','c:roundedCorners','c:style','c:clrMapOvr','c:pivotSource',
          'c:protection','c:chart','c:spPr','c:txPr','c:externalData','c:printSettings',
          'c:userShapes','c:extLst']
SCALING_ORDER=['c:logBase','c:orientation','c:max','c:min','c:extLst']

def place(parent, tag, order):
    ex=parent.find(qn(tag))
    if ex is not None: return ex
    el=etree.Element(qn(tag)); idx=order.index(tag)
    for ch in parent:
        key='c:'+etree.QName(ch).localname
        if key in order and order.index(key)>idx:
            ch.addprevious(el); return el
    parent.append(el); return el

def sub(el,tag,**a):
    e=etree.SubElement(el,qn(tag))
    for k,v in a.items(): e.set(k,str(v))
    return e

# ───────────────────────────────── workbook registry
# Every chart owns one worksheet whose layout is exactly the layout python-pptx
# writes into c:f — header row 1, categories in column A from row 2, series in
# columns B, C, D ... That makes the Excel link correct by construction and
# machine-verifiable (see verify_links.py).
CHARTS=[]          # (chart, sheet_name, cats, [(series_name, values)])

def sheet_name(base):
    nm=base[:31]; i=2
    used={c[1] for c in CHARTS}
    while nm in used: nm=(base[:28]+f"_{i}")[:31]; i+=1
    return nm

def write_workbook(path, rag_rows=None):
    wb=openpyxl.Workbook(); wb.remove(wb.active)
    idx=wb.create_sheet("Index")
    idx.append(["Sheet","Slide","Chart","Categories","Series"])
    for ch,nm,cats,series in CHARTS:
        ws=wb.create_sheet(nm)
        ws.append([None]+[n for n,_ in series])
        for i,c in enumerate(cats):
            ws.append([c]+[(v[i] if v[i] is not None else None) for _,v in series])
        ws.column_dimensions['A'].width=26
        idx.append([nm,nm.split('_')[0],nm,len(cats),len(series)])
    if rag_rows:
        r=wb.create_sheet("RAG_Rules")
        r.append(["Indicator","Basis","Green","Amber","Red","Value","Status"])
        for row in rag_rows: r.append(list(row))
        for col,w in zip("ABCDEFG",(34,30,18,20,18,16,10)):
            r.column_dimensions[col].width=w
    idx.column_dimensions['A'].width=26
    wb.save(path)

# ───────────────────────────────── chart helpers
def strip_chrome(chart, del_cat=True, del_val=True):
    cx=chart._chartSpace
    for ax in cx.iter(qn('c:catAx'), qn('c:valAx')):
        is_cat = ax.tag==qn('c:catAx')
        for g in ax.findall(qn('c:majorGridlines'))+ax.findall(qn('c:minorGridlines')): ax.remove(g)
        d=ax.find(qn('c:delete'))
        if d is None:
            d=etree.Element(qn('c:delete')); ax.find(qn('c:scaling')).addnext(d)
        d.set('val','1' if ((is_cat and del_cat) or ((not is_cat) and del_val)) else '0')
        for t in ('c:majorTickMark','c:minorTickMark'):
            e=ax.find(qn(t))
            if e is not None: e.set('val','none')
    for lg in cx.findall('.//'+qn('c:legend')): lg.getparent().remove(lg)
    chart.has_title=False

def transparent(chart):
    cx=chart._chartSpace
    pa=cx.find('.//'+qn('c:plotArea'))
    pasp=pa.find(qn('c:spPr'))
    if pasp is None:
        pasp=etree.Element(qn('c:spPr'))
        ext=pa.find(qn('c:extLst'))
        (ext.addprevious(pasp) if ext is not None else pa.append(pasp))
    for sp in (place(cx,'c:spPr',CS_ORDER), pasp):
        for c in list(sp): sp.remove(c)
        sub(sp,'a:noFill'); sub(sub(sp,'a:ln'),'a:noFill')

def plot_full(chart,x=0.0,y=0.0,w=1.0,h=1.0):
    pa=chart._chartSpace.find('.//'+qn('c:plotArea'))
    old=pa.find(qn('c:layout'))
    if old is not None: pa.remove(old)
    lay=etree.Element(qn('c:layout')); pa.insert(0,lay)
    ml=sub(lay,'c:manualLayout')
    sub(ml,'c:layoutTarget',val='inner'); sub(ml,'c:xMode',val='edge'); sub(ml,'c:yMode',val='edge')
    sub(ml,'c:x',val=x); sub(ml,'c:y',val=y); sub(ml,'c:w',val=w); sub(ml,'c:h',val=h)

def lock_text(chart, pt=T_CAP):
    cx=chart._chartSpace
    t=cx.find(qn('c:txPr'))
    if t is not None: cx.remove(t)
    t=place(cx,'c:txPr',CS_ORDER)
    sub(t,'a:bodyPr'); sub(t,'a:lstStyle')
    p=sub(t,'a:p'); sub(sub(p,'a:pPr'),'a:defRPr',sz=int(pt*100),b=1)
    sub(p,'a:endParaRPr',lang='en-US')

def nowrap(dl):
    tx=dl._element.find(qn('c:txPr'))
    if tx is None: return
    bp=tx.find(qn('a:bodyPr'))
    if bp is not None: bp.set('wrap','none'); bp.set('spcFirstLastPara','0')

def repoint(chart, sheet):
    for f in chart._chartSpace.iter(qn('c:f')):
        if f.text and '!' in f.text: f.text=sheet+'!'+f.text.split('!',1)[1]

CHART_PARTS=[]
def reg(gf, base, cats, series):
    """Register a chart, give it a private worksheet and point every c:f at it."""
    ch=gf.chart; CHART_PARTS.append(ch.part)
    nm=sheet_name(base)
    CHARTS.append((ch,nm,list(cats),[(n,list(v)) for n,v in series]))
    repoint(ch,nm)
    return ch

# ───────────────────────────────── drawing primitives
def text(sh,s,size,color,*,bold=False,caps=False,space=0,align=PP_ALIGN.LEFT,
         line=None,wrap=True,anchor=MSO_ANCHOR.TOP):
    tf=sh.text_frame; tf.word_wrap=wrap
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=anchor
    for i,ln in enumerate(str(s).split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        if line: p.line_spacing=line
        r=p.add_run(); r.text=ln.upper() if caps else ln
        r.font.size=Pt(size); r.font.bold=bold; r.font.name=FONT; r.font.color.rgb=color
        if space: r.font._rPr.set('spc',str(int(space*100)))
    return sh

def box(sp,x,y,w,h,fill=None,line=None,lw=0.75):
    s=sp.add_shape(MSO_SHAPE.RECTANGLE,I(x),I(y),I(w),I(h)); s.shadow.inherit=False
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb=fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb=line; s.line.width=Pt(lw)
    s.text_frame.text=""
    return s

def tbox(sp,x,y,w,h):
    t=sp.add_textbox(I(x),I(y),I(w),I(h)); t.text_frame.word_wrap=True; return t

TITLE_X=6.00; TITLE_W=8.00
def fit(size, txt, width, minsize=22, k=0.0088):
    """Shrink a single-line font size until the string fits `width` inches."""
    while size>minsize and len(txt)*size*k>width: size-=1
    return size

def title_block(sp, title, subtitle=None):
    ts=fit(T_TITLE,title,TITLE_W,24)
    text(tbox(sp,TITLE_X,TITLE_Y,TITLE_W,0.70),title,ts,NAVY,bold=True,
         align=PP_ALIGN.CENTER,wrap=False,anchor=MSO_ANCHOR.BOTTOM)
    if subtitle:
        ss=fit(T_SUB,subtitle,11.60,11,0.0100)
        text(tbox(sp,4.20,SUB_Y,11.60,0.28),subtitle,ss,MUTED,caps=True,space=1.2,
             align=PP_ALIGN.CENTER,wrap=False)
    box(sp,M,RULE_Y,RIGHT-M,0.024,fill=NAVY)

def marker(sp,x,y,label,color,note="",w=8.0):
    box(sp,x,y+0.05,0.17,0.17,fill=color)
    t=tbox(sp,x+0.32,y,min(w,RIGHT-(x+0.32)),0.30); tf=t.text_frame
    tf.margin_left=tf.margin_top=tf.margin_bottom=0; tf.word_wrap=False
    p=tf.paragraphs[0]
    r=p.add_run(); r.text=label.upper()
    r.font.size=Pt(T_SECT); r.font.bold=True; r.font.name=FONT; r.font.color.rgb=NAVY
    r.font._rPr.set('spc','130')
    if note:
        r2=p.add_run(); r2.text="    "+note
        r2.font.size=Pt(T_CAP); r2.font.bold=False; r2.font.name=FONT; r2.font.color.rgb=MUTED

def kpi_row(sp,y,items,h=1.26):
    n=len(items); w=(RIGHT-M-(n-1)*GUT)/n
    for i,(hd,val,cap,strip,vc) in enumerate(items):
        x=M+i*(w+GUT)
        box(sp,x,y,w,h,fill=WHITE,line=RULE); box(sp,x,y,w,0.42,fill=strip)
        text(tbox(sp,x+0.20,y+0.105,w-0.40,0.24),hd,T_CARD,WHITE,bold=True,caps=True,space=1.2)
        text(tbox(sp,x+0.20,y+0.54,w-0.40,0.46),val,T_BIG,vc,bold=True,wrap=False)
        if cap: text(tbox(sp,x+0.20,y+1.00,w-0.40,0.22),cap,T_CAP,MUTED,caps=True,space=0.9)
    return y+h

def panel(sp,x,y,w,h,title,strip=NAVY):
    box(sp,x,y,w,h,fill=WHITE,line=RULE)
    box(sp,x,y,w,0.42,fill=strip)
    text(tbox(sp,x+0.20,y+0.105,w-0.40,0.24),title,T_CARD,WHITE,bold=True,caps=True,space=1.2)
    return y+0.42

def rail_card(sp,x,y,w,h,title,rail):
    box(sp,x,y,w,h,fill=WHITE,line=RULE); box(sp,x,y,0.09,h,fill=rail)
    text(tbox(sp,x+0.26,y+0.16,w-0.46,0.24),title,T_CAP,MUTED,caps=True,space=0.9)
    return y+0.46

def rows(sp,x,y,w,items,*,rh=0.42,val_color=None,bold_val=True,sep=True,lsize=T_BODY,vsize=T_BODY):
    """label / value rows with hairline separators"""
    for i,(lab,val) in enumerate(items):
        yy=y+i*rh
        text(tbox(sp,x,yy+0.04,w*0.60,0.28),lab,lsize,BODY)
        vc=val_color(val) if callable(val_color) else (val_color or DEEP)
        t=tbox(sp,x+w*0.62,yy+0.04,w*0.38,0.28); t.text_frame.word_wrap=False
        text(t,val,vsize,vc,bold=bold_val,align=PP_ALIGN.RIGHT)
        if sep and i<len(items)-1: box(sp,x,yy+rh-0.015,w,0.008,fill=RULE)
    return y+len(items)*rh

def chip(sp,x,y,w,h,label,kind):
    fill={'g':GRNT,'a':AMBT,'r':REDT,'n':WASH}[kind]
    col={'g':GRNDK,'a':AMBDK,'r':REDDK,'n':BODY}[kind]
    box(sp,x,y,w,h,fill=fill)
    text(tbox(sp,x,y+(h-0.20)/2,w,0.22),label,T_CAP-1,col,bold=True,caps=True,
         space=0.8,align=PP_ALIGN.CENTER)

def milestone_rows(sp,x,y,w,items,rh=0.46):
    """(label, date, status, kind)"""
    for i,(lab,date,st,kind) in enumerate(items):
        yy=y+i*rh
        text(tbox(sp,x,yy+0.06,w*0.42,0.26),lab,T_BODY,BODY)
        t=tbox(sp,x+w*0.44,yy+0.06,w*0.21,0.26); t.text_frame.word_wrap=False
        text(t,date,T_BODY,DEEP,bold=True)
        if st: chip(sp,x+w*0.66,yy+0.03,w*0.34,0.30,st,kind)
        if i<len(items)-1: box(sp,x,yy+rh-0.015,w,0.008,fill=RULE)
    return y+len(items)*rh

def bullets(sp,x,y,w,h,lines,size=T_BODY,color=BODY):
    text(tbox(sp,x,y,w,h),"\n".join(lines),size,color,line=1.28)

# ───────────────────────────────── linked chart builders
def rail_chart(sp,x,y,w,label,val,scale,col,fmt="#,##0",sheet="Safety",sub_="",
               label_h=0.24,track_h=0.18,chart_h=0.40):
    text(tbox(sp,x,y,w*0.54,label_h),label,T_CAP,NAVY,bold=True,caps=True,space=0.8)
    if sub_:
        t=tbox(sp,x+w*0.55,y,w*0.45,label_h); text(t,sub_,T_CAP,MUTED,align=PP_ALIGN.RIGHT)
    ty=y+label_h+0.10
    box(sp,x,ty+ (chart_h-track_h)/2,w,track_h,fill=RULE)
    d=CategoryChartData(); d.categories=[label]; d.add_series(label,(val,))
    ch=reg(sp.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,I(x),I(ty),I(w),I(chart_h),d),
           sheet,[label],[(label,(val,))])
    strip_chrome(ch); transparent(ch); ch.has_legend=False; plot_full(ch)
    va=ch.value_axis; va.minimum_scale=0; va.maximum_scale=scale
    s=ch.series[0]
    s.format.fill.solid(); s.format.fill.fore_color.rgb=col; s.format.line.fill.background()
    ch.plots[0].gap_width=int((chart_h/track_h-1)*100)
    s.data_labels.show_value=True
    s.data_labels.font.size=Pt(T_CAP); s.data_labels.font.bold=True
    s.data_labels.font.name=FONT; s.data_labels.font.color.rgb=col
    s.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
    s.data_labels.number_format=fmt; s.data_labels.number_format_is_linked=False
    nowrap(s.data_labels); lock_text(ch)
    return ty+chart_h

def bar_chart(sp,x,y,w,h,cats,series,colors,sheet,*,maxv=None,fmt='#,##0',
              horizontal=True,cat_size=T_CAP,lbl_size=T_CAP,gap=70,stacked=False,
              label_pos=XL_LABEL_POSITION.OUTSIDE_END,label_all=False,plot=(0.0,0.02,1.0,0.96),
              point_colors=None,lbl_color=None):
    d=CategoryChartData(); d.categories=cats
    for nm,vals in series: d.add_series(nm,vals)
    ctype=(XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED) if horizontal \
          else (XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED)
    ch=reg(sp.add_chart(ctype,I(x),I(y),I(w),I(h),d),sheet,cats,series)
    strip_chrome(ch,del_cat=False,del_val=True); transparent(ch); ch.has_legend=False
    plot_full(ch,*plot)
    for s,c in zip(ch.series,colors):
        s.format.fill.solid(); s.format.fill.fore_color.rgb=c; s.format.line.fill.background()
    if point_colors:                      # RAG-colour individual data points
        for si,pcs in enumerate(point_colors):
            if pcs is None or si>=len(ch.series): continue
            for pi,pc in enumerate(pcs):
                if pc is None: continue
                pt=ch.series[si].points[pi]
                pt.format.fill.solid(); pt.format.fill.fore_color.rgb=pc
                pt.format.line.fill.background()
    ch.plots[0].gap_width=gap
    if stacked: ch.plots[0].overlap=100
    ca=ch.category_axis
    ca.tick_labels.font.size=Pt(cat_size); ca.tick_labels.font.bold=True
    ca.tick_labels.font.name=FONT; ca.tick_labels.font.color.rgb=NAVY
    ca.format.line.color.rgb=RULE
    if maxv:
        va=ch.value_axis; va.minimum_scale=0; va.maximum_scale=maxv
    if horizontal:
        sc=ca._element.find(qn('c:scaling'))
        place(sc,'c:orientation',SCALING_ORDER).set('val','maxMin')
    targets=ch.series if label_all else [ch.series[0]]
    for s in targets:
        s.data_labels.show_value=True
        s.data_labels.font.size=Pt(lbl_size); s.data_labels.font.bold=True
        s.data_labels.font.name=FONT
        s.data_labels.font.color.rgb=(lbl_color if lbl_color is not None else
            (WHITE if label_pos==XL_LABEL_POSITION.INSIDE_END else DEEP))
        s.data_labels.position=label_pos
        s.data_labels.number_format=fmt; s.data_labels.number_format_is_linked=False
        nowrap(s.data_labels)
    lock_text(ch,lbl_size)
    return ch

def ring(sp,x,y,size,pct,color,sheet,label,value_txt):
    d=CategoryChartData(); d.categories=["Achieved","Gap"]
    d.add_series(label,(pct,round(1-pct,4)))
    ch=reg(sp.add_chart(XL_CHART_TYPE.DOUGHNUT,I(x),I(y),I(size),I(size),d),
           sheet,["Achieved","Gap"],[(label,(pct,round(1-pct,4)))])
    ch.has_legend=False; ch.has_title=False; transparent(ch)
    ch.plots[0].has_data_labels=False; plot_full(ch,0.02,0.02,0.96,0.96)
    for p,c in zip(ch.plots[0].series[0].points,(color,RULE)):
        p.format.fill.solid(); p.format.fill.fore_color.rgb=c; p.format.line.fill.background()
    dp=ch._chartSpace.find('.//'+qn('c:doughnutChart'))
    for t in ('c:firstSliceAng','c:holeSize'):
        e=dp.find(qn(t))
        if e is not None: dp.remove(e)
    sub(dp,'c:firstSliceAng',val=0); sub(dp,'c:holeSize',val=78)
    text(tbox(sp,x,y+size/2-0.28,size,0.44),value_txt,T_BIG,color,bold=True,align=PP_ALIGN.CENTER)
    text(tbox(sp,x,y+size+0.04,size,0.24),label,T_CAP,MUTED,caps=True,space=1.0,align=PP_ALIGN.CENTER)

# ───────────────────────────────── table
def table(sp,x,y,w,headers,data,colw,*,rh=0.46,hh=0.46,cell=T_TBLC,hdr=T_TBLH,
          chips=None,zebra=True,aligns=None,cell_colors=None):
    total=sum(colw); colw=[c/total*w for c in colw]
    box(sp,x,y,w,hh,fill=NAVY)
    cx=x
    for i,hcell in enumerate(headers):
        th=tbox(sp,cx+0.12,y+0.05,colw[i]-0.24,hh-0.10)
        text(th,hcell,hdr,WHITE,bold=True,caps=True,space=0.8,
             align=(aligns[i] if aligns else PP_ALIGN.LEFT),line=1.05,
             anchor=MSO_ANCHOR.MIDDLE)
        cx+=colw[i]
    yy=y+hh
    for r,row in enumerate(data):
        hrow=rh(r) if callable(rh) else rh
        if zebra and r%2==1: box(sp,x,yy,w,hrow,fill=WASH)
        box(sp,x,yy+hrow-0.008,w,0.008,fill=RULE)
        cx=x
        for i,val in enumerate(row):
            al=(aligns[i] if aligns else PP_ALIGN.LEFT)
            if chips and (r,i) in chips:
                chip(sp,cx+0.12,yy+(hrow-0.30)/2,colw[i]-0.24,0.30,str(val),chips[(r,i)])
            else:
                col=(cell_colors or {}).get((r,i), BODY if i>0 else DEEP)
                text(tbox(sp,cx+0.12,yy+0.10,colw[i]-0.24,hrow-0.18),str(val),cell,
                     col,bold=(i==0 or (r,i) in (cell_colors or {})),align=al,line=1.18)
            cx+=colw[i]
        yy+=hrow
    box(sp,x,y,w,yy-y,fill=None,line=RULE)
    return yy

def rag_legend(sp, y=10.70):
    """Compact RAG key, bottom-right, identical on every data slide."""
    items=(("On track",GREEN),("Watch",AMBER),("Critical",RED))
    wds=[1.42,1.10,1.28]; tot=sum(wds)+0.90
    x=RIGHT-tot
    text(tbox(sp,x-1.28,y,1.20,0.24),"RAG KEY",T_CAP-2,MUTED,caps=True,space=1.0,
         align=PP_ALIGN.RIGHT)
    for (lab,c),wd in zip(items,wds):
        box(sp,x,y+0.045,0.17,0.15,fill=c)
        text(tbox(sp,x+0.26,y,wd,0.24),lab,T_CAP-2,BODY,caps=True,space=0.8)
        x+=wd+0.10

# ───────────────────────────────── slide utilities
def is_logo(sh):
    """True only for the three partner logos in the top band and the footer strip."""
    E=914400
    if sh.left is None: return False
    l,t,w,h=sh.left/E,sh.top/E,sh.width/E,sh.height/E
    if t>10.75 and h<0.70: return True                       # footer colour strip
    has_blip = sh._element.find('.//'+qn('a:blip')) is not None
    if t<1.35 and h<1.25 and w<=4.20 and has_blip: return True
    return False

def clear_slide(slide, keep_pictures=False, keep_ole=False):
    """Strip the slide back to its logos (de-duplicated) plus any retained artwork."""
    E=914400; kept=[]; logos=[]
    for sh in list(slide.shapes):
        keep=False; logo=False
        if is_logo(sh): keep=True; logo=True
        if keep_pictures and sh.shape_type==MSO_SHAPE_TYPE.PICTURE: keep=True
        if keep_ole and sh.shape_type==MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT: keep=True
        if keep and logo:
            g=(sh.width/E, sh.height/E, sh.left/E, sh.top/E)
            if any(abs(g[0]-o[0])<0.06 and abs(g[1]-o[1])<0.06 and
                   abs(g[2]-o[2])<0.40 and abs(g[3]-o[3])<0.40 for o in logos):
                keep=False                      # duplicate logo inherited from the original
            else:
                logos.append(g)
        if keep: kept.append(sh)
        else: sh._element.getparent().remove(sh._element)
    return kept

import hashlib, collections
CHROME=set()          # md5s of template/decorative artwork, filled by scan_chrome()

def scan_chrome(prs, lo=5, hi=38, threshold=3):
    """Any image reused on >=threshold slides is template chrome, not a photograph."""
    seen=collections.defaultdict(set)
    for n in range(lo,hi+1):
        sl=prs.slides[n-1]
        for b in sl._element.iter(qn('a:blip')):
            rid=b.get(qn('r:embed')) or b.get(qn('r:link'))
            if not rid: continue
            try: img=sl.part.related_part(rid)
            except Exception: continue
            seen[hashlib.md5(img.blob).hexdigest()].add(n)
    CHROME.clear()
    CHROME.update(m for m,s in seen.items() if len(s)>=threshold)
    return CHROME

def harvest_images(slide, min_bytes=3000):
    """Photographs on the slide, in document order, de-duplicated, chrome removed."""
    out=[]; seen=set()
    for b in slide._element.iter(qn('a:blip')):
        rid=b.get(qn('r:embed')) or b.get(qn('r:link'))
        if not rid: continue
        try: img=slide.part.related_part(rid)
        except Exception: continue
        blob=img.blob
        if len(blob)<min_bytes: continue
        m=hashlib.md5(blob).hexdigest()
        if m in CHROME or m in seen: continue
        seen.add(m)
        out.append({'blob':blob,'md5':m})
    return out

def place_photo(sp, blob, x, y, w, h):
    """Insert image cropped to exactly fill the x/y/w/h cell."""
    from PIL import Image
    try:
        im=Image.open(io.BytesIO(blob)); iw,ih=im.size
    except Exception:
        iw,ih=(4,3)
    src=iw/ih; dst=w/h
    pic=sp.add_picture(io.BytesIO(blob),I(x),I(y),I(w),I(h))
    if src>dst:                      # too wide -> crop sides
        keep=dst/src; c=(1-keep)/2
        pic.crop_left=c; pic.crop_right=c
    elif src<dst:                    # too tall -> crop top/bottom
        keep=src/dst; c=(1-keep)/2
        pic.crop_top=c; pic.crop_bottom=c
    pic.line.color.rgb=RULE; pic.line.width=Pt(0.75)
    return pic

def photo_slide(sp, images, captions, cols=None):
    n=len(images)
    if n==0: return
    cols=cols or (3 if n>=5 else (2 if n>=2 else 1))
    rowsn=(n+cols-1)//cols
    availw=RIGHT-M; availh=BODY_BOT-BODY_TOP
    caph=0.34 if captions else 0.0
    cw=(availw-(cols-1)*GUT)/cols
    chh=(availh-(rowsn-1)*GUT)/rowsn
    ph=chh-caph
    for i,img in enumerate(images[:cols*rowsn]):
        r,c=divmod(i,cols)
        x=M+c*(cw+GUT); y=BODY_TOP+r*(chh+GUT)
        place_photo(sp,img['blob'],x,y,cw,ph)
        if captions:
            cap=captions[i] if i<len(captions) else ""
            if cap:
                text(tbox(sp,x,y+ph+0.06,cw,caph-0.06),cap,T_CAP,BODY,line=1.1)
