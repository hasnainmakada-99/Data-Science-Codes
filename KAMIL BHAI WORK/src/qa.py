import sys, zipfile, hashlib, collections, re
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE_TYPE as T
F=sys.argv[1]; E=914400; PW,PH=20.0,11.25
prs=Presentation(F); fails=[]
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
n=len(prs.slides._sldIdLst)
if n!=38: fails.append(f"slide count {n}")
# off-canvas + overlaps
def rects(sl):
    out=[]
    for sh in sl.shapes:
        if sh.left is None: continue
        out.append((sh, sh.left/E, sh.top/E, sh.width/E, sh.height/E))
    return out
off=0; ovl=0
for i,sl in enumerate(prs.slides,1):
    if i<5: continue
    rs=rects(sl)
    for sh,l,t,w,h in rs:
        if l<-0.5 or t<-0.5 or l+w>PW+0.5 or t+h>PH+0.5:
            off+=1; print(f"  OFFCANVAS s{i}: {sh.shape_type} {l:.2f},{t:.2f} {w:.2f}x{h:.2f}")
    tb=[(sh,l,t,w,h) for sh,l,t,w,h in rs
        if sh.shape_type==T.TEXT_BOX and sh.has_text_frame and sh.text_frame.text.strip()]
    for a in range(len(tb)):
        for b in range(a+1,len(tb)):
            _,l1,t1,w1,h1=tb[a]; _,l2,t2,w2,h2=tb[b]
            ox=min(l1+w1,l2+w2)-max(l1,l2); oy=min(t1+h1,t2+h2)-max(t1,t2)
            if ox>0.05 and oy>0.05:
                ovl+=1
                print(f"  OVERLAP s{i}: '{tb[a][0].text_frame.text[:22]}' / '{tb[b][0].text_frame.text[:22]}' {ox:.2f}x{oy:.2f}")
if off: fails.append(f"{off} off-canvas")
if ovl: fails.append(f"{ovl} text overlaps")
# duplicate ids
for i,sl in enumerate(prs.slides,1):
    ids=[e.get('id') for e in sl._element.iter(qn('p:cNvPr'))]
    d=[k for k,v in collections.Counter(ids).items() if v>1]
    if d: fails.append(f"dup ids slide {i}: {d}")
# package level
z=zipfile.ZipFile(F)
ct=z.read('[Content_Types].xml').decode()
charts=[n for n in z.namelist() if re.match(r'ppt/charts/chart\d+\.xml$',n)]
miss=[c for c in charts if '/'+c not in ct]
if miss: fails.append(f"chart content-type missing: {miss}")
bad=[c for c in charts if 'Sheet1!' in z.read(c).decode('utf8','ignore')]
if bad: fails.append(f"Sheet1! refs in {bad}")
# unresolved r:id
for name in z.namelist():
    if not name.startswith('ppt/slides/slide') or not name.endswith('.xml'): continue
    rels=name.replace('slides/','slides/_rels/')+'.rels'
    have=set(re.findall(r'Id="([^"]+)"', z.read(rels).decode()))
    used=set(re.findall(r'r:(?:id|embed|link)="([^"]+)"', z.read(name).decode()))
    if used-have: fails.append(f"{name}: unresolved {used-have}")
print("MASTER md5:", hashlib.md5(open('master.pptx','rb').read()).hexdigest())
print("charts:", len(charts))
print("FAIL: "+"; ".join(fails) if fails else "QA PASS")
